import io
import secrets
from datetime import datetime, timezone

import openpyxl
from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile, status
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.deps import OVERSIGHT_ROLES, get_current_user, require_role
from app.core.i18n import get_locale, t
from app.core.security import hash_password, verify_password
from app.db.session import get_db
from app.models.department import Department
from app.models.position import Position
from app.models.user import AcademicDegree, User, UserRole
from app.models.kpi_template import KpiTemplate
from app.schemas.auth import PasswordChange, UserOut, UserPreferencesUpdate
from app.schemas.employee import (
    PRIVILEGED_ROLES,
    EmployeeCreate,
    EmployeeUpdate,
    ImportRowResult,
    ImportSummary,
    RestrictRequest,
)
from app.services import hemis_rest
from app.services.audit import log_action

router = APIRouter(prefix="/users", tags=["users"])

IMPORT_REQUIRED_COLUMNS = {"email", "full_name", "role"}


def _department_and_descendant_ids(db: Session, department_id: int) -> list[int]:
    """A department plus every descendant (e.g. a faculty plus all of its kafedras)."""
    all_departments = db.scalars(select(Department)).all()
    children_by_parent: dict[int, list[int]] = {}
    for d in all_departments:
        if d.parent_department_id is not None:
            children_by_parent.setdefault(d.parent_department_id, []).append(d.id)

    ids = [department_id]
    frontier = [department_id]
    while frontier:
        current = frontier.pop()
        for child_id in children_by_parent.get(current, []):
            if child_id not in ids:
                ids.append(child_id)
                frontier.append(child_id)
    return ids


@router.get("", response_model=list[UserOut])
def list_users(
    faculty_id: int | None = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> list[UserOut]:
    if current_user.role in OVERSIGHT_ROLES:
        query = select(User)
        if faculty_id is not None:
            department_ids = _department_and_descendant_ids(db, faculty_id)
            query = query.where(User.department_id.in_(department_ids))
        users = db.scalars(query).all()
    elif current_user.role == UserRole.manager:
        users = db.scalars(
            select(User).where((User.manager_id == current_user.id) | (User.id == current_user.id))
        ).all()
    else:
        users = [current_user]

    return [UserOut.model_validate(u) for u in users]


@router.patch("/me", response_model=UserOut)
def update_me(
    payload: UserPreferencesUpdate,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> UserOut:
    if payload.full_name is not None:
        current_user.full_name = payload.full_name
    if payload.theme_preference is not None:
        current_user.theme_preference = payload.theme_preference
    if payload.sidebar_collapsed is not None:
        current_user.sidebar_collapsed = payload.sidebar_collapsed
    if payload.full_name is not None:
        log_action(db, current_user.id, "update", "user", current_user.id)
    db.commit()
    db.refresh(current_user)
    return UserOut.model_validate(current_user)


@router.patch("/me/password", status_code=status.HTTP_204_NO_CONTENT)
def change_own_password(
    payload: PasswordChange,
    request: Request,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> None:
    locale = get_locale(request)
    if not verify_password(payload.current_password, current_user.password_hash):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=t("incorrect_current_password", locale))

    current_user.password_hash = hash_password(payload.new_password)
    log_action(db, current_user.id, "update", "user", current_user.id)
    db.commit()


def _validate_assignment(
    db: Session,
    locale: str,
    department_id: int | None,
    position_id: int | None,
    manager_id: int | None,
    kpi_template_id: int | None = None,
) -> None:
    if department_id is not None and db.get(Department, department_id) is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=t("department_not_found", locale))

    if position_id is not None:
        position = db.get(Position, position_id)
        if position is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=t("position_not_found", locale))
        if department_id is not None and position.department_id != department_id:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST, detail=t("position_department_mismatch", locale)
            )

    if manager_id is not None and db.get(User, manager_id) is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=t("manager_not_found", locale))

    if kpi_template_id is not None and db.get(KpiTemplate, kpi_template_id) is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=t("kpi_template_not_found", locale))


@router.post("", response_model=UserOut, status_code=status.HTTP_201_CREATED)
def create_employee(
    payload: EmployeeCreate,
    request: Request,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role("super_admin", "admin")),
) -> UserOut:
    locale = get_locale(request)

    if payload.role in PRIVILEGED_ROLES and current_user.role != UserRole.super_admin:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail=t("not_enough_permissions", locale))

    if db.scalar(select(User).where(User.email == payload.email)) is not None:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=t("email_already_exists", locale))

    _validate_assignment(db, locale, payload.department_id, payload.position_id, payload.manager_id, payload.kpi_template_id)

    user = User(
        email=payload.email,
        password_hash=hash_password(payload.password),
        full_name=payload.full_name,
        role=UserRole(payload.role),
        department_id=payload.department_id,
        position_id=payload.position_id,
        manager_id=payload.manager_id,
        kpi_template_id=payload.kpi_template_id,
        academic_degree=AcademicDegree(payload.academic_degree),
        bonus_fund_override=payload.bonus_fund_override,
    )
    db.add(user)
    db.flush()
    log_action(db, current_user.id, "create", "user", user.id)
    db.commit()
    db.refresh(user)
    return UserOut.model_validate(user)


@router.get("/import-template")
def download_import_template(
    current_user: User = Depends(require_role("super_admin", "admin")),
) -> StreamingResponse:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Xodimlar"
    sheet.append(["email", "full_name", "role", "department", "position", "manager_email"])
    sheet.append(["aziza@example.com", "Aziza Karimova", "employee", "IT bo'limi", "Dasturchi", "manager@example.com"])

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=xodimlar_shablon.xlsx"},
    )


@router.post("/import", response_model=ImportSummary)
async def import_employees(
    request: Request,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role("super_admin", "admin")),
) -> ImportSummary:
    locale = get_locale(request)

    raw = await file.read()
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        sheet = workbook.active
        sheet_rows = sheet.iter_rows(values_only=True)
        header_row = next(sheet_rows)
    except Exception:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=t("excel_unreadable", locale))

    header_index = {
        str(name).strip().lower(): idx for idx, name in enumerate(header_row) if name is not None
    }
    if not IMPORT_REQUIRED_COLUMNS.issubset(header_index.keys()):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=t("excel_missing_columns", locale))

    def cell(row: tuple, column: str) -> str:
        idx = header_index.get(column)
        if idx is None or idx >= len(row) or row[idx] is None:
            return ""
        return str(row[idx]).strip()

    departments_by_name = {d.name.strip().lower(): d for d in db.scalars(select(Department)).all()}
    positions_by_key = {
        (p.department_id, p.title.strip().lower()): p for p in db.scalars(select(Position)).all()
    }

    results: list[ImportRowResult] = []
    created_count = 0

    for index, raw_row in enumerate(sheet_rows, start=2):
        if raw_row is None or all(value is None for value in raw_row):
            continue
        row = {
            "email": cell(raw_row, "email"),
            "full_name": cell(raw_row, "full_name"),
            "role": cell(raw_row, "role"),
            "department": cell(raw_row, "department"),
            "position": cell(raw_row, "position"),
            "manager_email": cell(raw_row, "manager_email"),
        }
        email = (row.get("email") or "").strip().lower()
        full_name = (row.get("full_name") or "").strip()
        role_raw = (row.get("role") or "").strip().lower()
        department_name = (row.get("department") or "").strip()
        position_title = (row.get("position") or "").strip()
        manager_email = (row.get("manager_email") or "").strip().lower()

        if not email or not full_name:
            results.append(ImportRowResult(row=index, email=email, status="error", error=t("row_missing_required_fields", locale)))
            continue

        if role_raw not in ("manager", "employee"):
            results.append(ImportRowResult(row=index, email=email, status="error", error=t("invalid_role", locale)))
            continue

        if db.scalar(select(User).where(User.email == email)) is not None:
            results.append(
                ImportRowResult(row=index, email=email, status="error", error=t("email_already_exists", locale))
            )
            continue

        department_id = None
        if department_name:
            department = departments_by_name.get(department_name.lower())
            if department is None:
                results.append(
                    ImportRowResult(row=index, email=email, status="error", error=t("department_not_found", locale))
                )
                continue
            department_id = department.id

        position_id = None
        if position_title:
            position = positions_by_key.get((department_id, position_title.lower())) if department_id else None
            if position is None:
                results.append(
                    ImportRowResult(row=index, email=email, status="error", error=t("position_not_found", locale))
                )
                continue
            position_id = position.id

        manager_id = None
        if manager_email:
            manager = db.scalar(select(User).where(User.email == manager_email))
            if manager is None:
                results.append(
                    ImportRowResult(row=index, email=email, status="error", error=t("manager_not_found", locale))
                )
                continue
            manager_id = manager.id

        temporary_password = secrets.token_urlsafe(9)
        user = User(
            email=email,
            password_hash=hash_password(temporary_password),
            full_name=full_name,
            role=UserRole(role_raw),
            department_id=department_id,
            position_id=position_id,
            manager_id=manager_id,
        )
        db.add(user)
        db.flush()
        log_action(db, current_user.id, "create", "user", user.id)

        created_count += 1
        results.append(
            ImportRowResult(row=index, email=email, status="created", temporary_password=temporary_password)
        )

    db.commit()

    return ImportSummary(total=len(results), created=created_count, failed=len(results) - created_count, results=results)


@router.patch("/{user_id}", response_model=UserOut)
def update_employee(
    user_id: int,
    payload: EmployeeUpdate,
    request: Request,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role("super_admin", "admin")),
) -> UserOut:
    locale = get_locale(request)
    user = db.get(User, user_id)
    if user is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=t("user_not_found", locale))

    target_is_privileged = user.role.value in PRIVILEGED_ROLES or user.role == UserRole.super_admin
    new_role_is_privileged = payload.role is not None and payload.role in PRIVILEGED_ROLES
    if (target_is_privileged or new_role_is_privileged) and current_user.role != UserRole.super_admin:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail=t("not_enough_permissions", locale))

    _validate_assignment(db, locale, payload.department_id, payload.position_id, payload.manager_id, payload.kpi_template_id)

    fields_set = payload.model_fields_set
    if payload.full_name is not None:
        user.full_name = payload.full_name
    if payload.role is not None:
        user.role = UserRole(payload.role)
    if "department_id" in fields_set:
        user.department_id = payload.department_id
    if "position_id" in fields_set:
        user.position_id = payload.position_id
    if "manager_id" in fields_set:
        user.manager_id = payload.manager_id
    if "kpi_template_id" in fields_set:
        user.kpi_template_id = payload.kpi_template_id
    if payload.academic_degree is not None:
        user.academic_degree = AcademicDegree(payload.academic_degree)
    if payload.is_active is not None:
        user.is_active = payload.is_active
    if "bonus_fund_override" in fields_set:
        user.bonus_fund_override = payload.bonus_fund_override

    log_action(db, current_user.id, "update", "user", user.id)
    db.commit()
    db.refresh(user)
    return UserOut.model_validate(user)


@router.get("/removed", response_model=list[UserOut])
def list_removed_users(
    current_user: User = Depends(require_role("super_admin", "admin")),
    db: Session = Depends(get_db),
) -> list[UserOut]:
    users = db.scalars(select(User).where(User.is_active.is_(False))).all()
    return [UserOut.model_validate(u) for u in users]


@router.get("/{user_id}", response_model=UserOut)
def get_employee(
    user_id: int,
    request: Request,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> UserOut:
    locale = get_locale(request)
    user = db.get(User, user_id)
    if user is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=t("user_not_found", locale))
    is_self = current_user.id == user.id
    is_admin = current_user.role in OVERSIGHT_ROLES
    is_own_report = current_user.role == UserRole.manager and user.manager_id == current_user.id
    if not (is_self or is_admin or is_own_report):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail=t("not_enough_permissions", locale))
    return UserOut.model_validate(user)


@router.post("/{user_id}/hemis-sync", response_model=UserOut)
def sync_employee_from_hemis(
    user_id: int,
    request: Request,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role("super_admin", "admin")),
) -> UserOut:
    """Pulls fresh data for this one employee from HEMIS's server-side REST
    directory (separate token from the OAuth login flow - see
    app/services/hemis_rest.py). Works for any employee with a known
    `hemis_employee_id_number`, regardless of whether they've logged in via
    HEMIS themselves recently."""
    locale = get_locale(request)
    user = db.get(User, user_id)
    if user is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=t("user_not_found", locale))
    if not user.hemis_employee_id_number:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=t("hemis_sync_no_identifier", locale))

    try:
        meta = hemis_rest.fetch_employee_by_id_number(user.hemis_employee_id_number)
    except hemis_rest.HemisRestError as exc:
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail=str(exc)) from exc
    if meta is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=t("hemis_sync_not_found", locale))

    hemis_rest.apply_employee_meta(user, meta)
    user.hemis_rest_synced_at = datetime.now(timezone.utc)

    log_action(db, current_user.id, "hemis_sync", "user", user.id)
    db.commit()
    db.refresh(user)
    return UserOut.model_validate(user)


@router.patch("/{user_id}/restrict", response_model=UserOut)
def set_restriction(
    user_id: int,
    payload: RestrictRequest,
    request: Request,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role("super_admin", "admin")),
) -> UserOut:
    locale = get_locale(request)
    user = db.get(User, user_id)
    if user is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=t("user_not_found", locale))

    user.is_restricted = payload.is_restricted
    log_action(db, current_user.id, "restrict" if payload.is_restricted else "unrestrict", "user", user.id)
    db.commit()
    db.refresh(user)
    return UserOut.model_validate(user)


@router.patch("/{user_id}/restore", response_model=UserOut)
def restore_user(
    user_id: int,
    request: Request,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role("super_admin", "admin")),
) -> UserOut:
    locale = get_locale(request)
    user = db.get(User, user_id)
    if user is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=t("user_not_found", locale))

    user.is_active = True
    log_action(db, current_user.id, "restore", "user", user.id)
    db.commit()
    db.refresh(user)
    return UserOut.model_validate(user)
